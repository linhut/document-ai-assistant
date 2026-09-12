# (c) 2026 Jose AI (https://www.linhut.cn)
# https://github.com/linhut/document-ai-assistant
# Licensed under the MIT License. See the LICENSE file for details.

"""PDF 双输出服务单元测试（原创实现，转换器缺失时自动跳过实转用例）。"""

import pytest
from docx import Document
from pathlib import Path

from services.pdf_export import convert_docx_to_pdf, find_libreoffice, has_converter

HAS_CONV = has_converter()


def _dummy_docx(tmp_path: Path) -> Path:
    docx = tmp_path / "样例公文.docx"
    d = Document()
    d.add_paragraph("关于测试 PDF 导出的通知")
    d.add_paragraph("正文内容示例。")
    d.save(docx)
    return docx


def test_find_libreoffice_returns_path_or_none():
    result = find_libreoffice()
    assert result is None or (isinstance(result, str) and Path(result).exists())


def test_has_converter_boolean():
    assert isinstance(has_converter(), bool)


def test_missing_source_raises():
    with pytest.raises(RuntimeError, match="源文件不存在"):
        convert_docx_to_pdf(Path("不存在.docx"))


@pytest.mark.skipif(HAS_CONV, reason="转换器可用时跳过降级路径测试")
def test_no_converter_raises_actionable_error(tmp_path):
    docx = _dummy_docx(tmp_path)
    with pytest.raises(RuntimeError) as exc:
        convert_docx_to_pdf(docx, tmp_path / "out")
    msg = str(exc.value)
    assert "LibreOffice" in msg or "docx2pdf" in msg


@pytest.mark.skipif(not HAS_CONV, reason="无转换器，跳过实转测试")
def test_convert_produces_pdf(tmp_path):
    docx = _dummy_docx(tmp_path)
    out = tmp_path / "out"
    pdf = convert_docx_to_pdf(docx, out)
    assert pdf.exists()
    assert pdf.suffix.lower() == ".pdf"
    assert pdf.stat().st_size > 100